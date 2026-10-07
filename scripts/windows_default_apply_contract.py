"""Real installed console apply, with an isolated producer for each business step."""

import hashlib
import importlib.metadata
import json
import os
import shutil
import subprocess
import sys
from pathlib import Path

from windows_apply_fixture import hosted_scope


def authorization(root):
    value = json.loads((root / "authorization.json").read_text(encoding="utf-8"))
    if (
        value["root"] != str(root)
        or value["sha"] != os.environ["GITHUB_SHA"]
        or (root / ".mara-pytest-owner").read_text() != value["owner"]
    ):
        raise RuntimeError("Default MARA scope does not match its authorization")
    return value


def installed_wheels():
    result = {}
    for name in ("kotaemon", "ktem", "mara-research-cli", "mara-app"):
        dist = importlib.metadata.distribution(name)
        metadata = dist.read_text("direct_url.json")
        assert metadata is not None
        direct = json.loads(metadata)
        assert "archive_info" in direct and not direct.get("dir_info", {}).get(
            "editable"
        )
        result[name] = {"version": dist.version, "wheel": direct}
    return result


def prepare(root):
    from pptx import Presentation
    from slide_cli import deck, paths
    from slide_cli.session_store import SlideSessionStore

    assert Path(deck.__file__).resolve().is_relative_to(Path(sys.prefix).resolve())
    assert paths.get_slide_runtime_paths().data_dir == root
    cwd = root / "fixture" / "outside 中文"
    cwd.mkdir(parents=True)
    source = cwd / "原始.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "Owned original"
    slide.placeholders[1].text = "Owned body"
    presentation.save(source)
    snapshot = deck.load_deck_snapshot(source)
    shape, body = snapshot.slides[0].shapes[:2]
    patch = deck.DeckPatch(
        "Owned patch",
        [
            deck.TextReplaceOp(1, shape.target_id, shape.text, "Owned edited"),
            deck.TextReplaceOp(1, body.target_id, "stale value", "Must not apply"),
        ],
    )
    store = SlideSessionStore()
    commands = []
    for index, command in enumerate(("MARA", "MARA-cli")):
        session = store.create_session(
            mode="run",
            input_path=str(source),
            cwd=str(cwd),
            session_id=f"owned-{index}",
        )
        store.append_event(
            session.session_id, {"kind": "final", "patch": patch.as_dict()}
        )
        commands.append(
            {
                "console": command,
                "session": session.session_id,
                "output": str(cwd / f"修改后-{index}.pptx"),
                "expected": 0,
            }
        )
    commands.append(
        {
            "console": "MARA",
            "session": "not-a-session",
            "output": str(cwd / "missing-result.pptx"),
            "expected": 1,
        }
    )
    (root / "plan.json").write_text(
        json.dumps(
            {
                "commands": commands,
                "cwd": str(cwd),
                "source": str(source),
                "source_sha256": hashlib.sha256(source.read_bytes()).hexdigest(),
                "module_path": deck.__file__,
                "installed_wheels": installed_wheels(),
            },
            ensure_ascii=False,
            indent=2,
        ),
        encoding="utf-8",
    )


def verify(root):
    import psutil
    from slide_cli import deck, paths
    from slide_cli.session_store import SlideSessionStore

    assert paths.get_slide_runtime_paths().data_dir == root
    plan = json.loads((root / "plan.json").read_text(encoding="utf-8"))
    receipts = json.loads((root / "console-results.json").read_text(encoding="utf-8"))
    store = SlideSessionStore()
    for item, receipt in zip(plan["commands"], receipts):
        assert receipt["exit_code"] == item["expected"]
        output = Path(item["output"])
        if item["expected"]:
            assert not output.exists()
            continue
        payload = receipt["payload"]
        assert payload["applied_count"] == 1 and payload["skipped_count"] == 1
        assert payload["session_id"] == item["session"]
        reloaded = deck.load_deck_snapshot(output)
        assert reloaded.slides[0].shapes[0].text == "Owned edited"
        assert reloaded.slides[0].shapes[1].text == "Owned body"
        saved = store.load_session(item["session"])
        assert saved is not None and saved.status == "completed"
        assert saved.output_path == str(output)
        assert [event["kind"] for event in saved.events] == ["final", "apply"]
    assert len(receipts) == len(plan["commands"]) == 3
    assert (
        hashlib.sha256(Path(plan["source"]).read_bytes()).hexdigest()
        == plan["source_sha256"]
    )
    assert not list(Path(plan["cwd"]).glob(".mara-deck-*"))
    guards = [json.loads(p.read_text()) for p in root.glob("console-guard-*.json")]
    assert len(guards) == 3 and all(row["closed"] for row in guards)
    assert all(not psutil.pid_exists(row["pid"]) for row in guards)
    (root / "result.json").write_text(
        json.dumps(
            {
                **plan,
                "commands": receipts,
                "console_processes": guards,
                "default_paths": True,
                "PYTHONPATH": False,
            },
            ensure_ascii=False,
            indent=2,
        ),
        encoding="utf-8",
    )


def business_step(mode):
    root, identity = hosted_scope()
    os.environ["MARA_DIAGNOSTIC_ISOLATION_REQUIRED"] = "1"
    from pytest_runtime_isolation import start_process_test_runtime

    if mode == "prepare":
        if root.exists():
            raise RuntimeError("Default MARA scope already exists; refusing reuse")
        # The actual Known Folder was resolved before this first write.
        root.mkdir(parents=True)
        runtime = start_process_test_runtime(evidence_dir=root)
        value = {
            **identity,
            "root": str(root),
            "sha": os.environ["GITHUB_SHA"],
            "run": os.environ["GITHUB_RUN_ID"],
            "owner": runtime.owner_token,
            "fresh_before_write": True,
        }
        (root / "authorization.json").write_text(json.dumps(value, indent=2))
        print(json.dumps({"authorized_scope": str(root)}), flush=True)
    else:
        authorization(root)
        runtime = start_process_test_runtime()
    receipt = root / f"producer-{mode}.json"
    receipt.write_text(json.dumps({"pid": os.getpid(), "closed": False}))
    try:
        assert "PYTHONPATH" not in os.environ
        (prepare if mode == "prepare" else verify)(root)
        assert runtime.process_guard is not None and not runtime.process_guard.denials
    finally:
        runtime.close()
        receipt.write_text(json.dumps({"pid": os.getpid(), "closed": True}))
        (root / f"{mode}-cleanup.json").write_text(
            json.dumps(
                {
                    "runtime_removed": not runtime.paths.root.exists(),
                    "pid": os.getpid(),
                }
            )
        )


def run_console(root, item, environment):
    executable = Path(sys.prefix) / "Scripts" / f"{item['console']}.exe"
    argv = [
        str(executable),
        "apply",
        item["session"],
        "--output",
        item["output"],
        "--json",
    ]
    completed = subprocess.run(
        argv,
        cwd=Path(item["output"]).parent,
        env=environment,
        capture_output=True,
        text=True,
        encoding="utf-8",
        timeout=60,
    )
    label = f"{item['console']}-{item['session']}"
    (root / f"{label}.stdout").write_text(completed.stdout, encoding="utf-8")
    (root / f"{label}.stderr").write_text(completed.stderr, encoding="utf-8")
    assert completed.returncode == item["expected"], (
        label,
        completed.returncode,
        completed.stderr,
    )
    return {
        "command": argv,
        "exit_code": completed.returncode,
        "payload": json.loads(completed.stdout) if item["expected"] == 0 else None,
    }


def coordinate():
    root, _ = hosted_scope()
    if root.exists():
        raise RuntimeError("Default MARA scope already exists; refusing reuse")
    assert "PYTHONPATH" not in os.environ
    script = str(Path(__file__).resolve())
    # This controller imports no business module. Every producer starts with the
    # disposable account's original profile and installs the original guard.
    subprocess.run([sys.executable, "-B", script, "--prepare"], check=True, timeout=90)
    value = authorization(root)
    environment = dict(
        os.environ,
        MARA_DIAGNOSTIC_ISOLATION_REQUIRED="1",
        MARA_DIAGNOSTIC_EVIDENCE_DIR=str(root),
        MARA_DIAGNOSTIC_EVIDENCE_OWNER=value["owner"],
    )
    plan = json.loads((root / "plan.json").read_text(encoding="utf-8"))
    console_env = dict(environment, MARA_CI_DEFAULT_APPLY_CHILD="1")
    receipts = [run_console(root, item, console_env) for item in plan["commands"]]
    (root / "console-results.json").write_text(
        json.dumps(receipts, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    subprocess.run(
        [sys.executable, "-B", script, "--verify"],
        env=environment,
        check=True,
        timeout=90,
    )


def collect(destination):
    import psutil

    root, _ = hosted_scope()
    authorization(root)
    destination = Path(destination).resolve()
    runner_temp = Path(os.environ["RUNNER_TEMP"]).resolve()
    if not destination.is_relative_to(runner_temp) or destination.exists():
        raise RuntimeError("Evidence destination must be fresh and within RUNNER_TEMP")
    shutil.copytree(root, destination)
    # Do not remove a scope whose application producer is still running, even
    # after its console launcher has returned or a subprocess timeout fired.
    for pattern in ("producer-*.json", "console-guard-*.json"):
        for receipt in root.glob(pattern):
            producer = json.loads(receipt.read_text())
            if not producer["closed"] or psutil.pid_exists(producer["pid"]):
                raise RuntimeError("Owned application producer has not finished")
    shutil.rmtree(root)
    (destination / "scope-cleanup.json").write_text(
        json.dumps({"default_scope_removed": not root.exists()})
    )


if __name__ == "__main__":
    if len(sys.argv) == 3 and sys.argv[1] == "--collect":
        collect(sys.argv[2])
    elif len(sys.argv) == 2 and sys.argv[1] in {"--prepare", "--verify"}:
        business_step(sys.argv[1][2:])
    elif len(sys.argv) == 1:
        coordinate()
    else:
        raise SystemExit(
            "Expected no arguments, --prepare, --verify or --collect DESTINATION"
        )
