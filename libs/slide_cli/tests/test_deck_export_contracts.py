"""Conversion boundaries and real deck edit/reload identities."""

import subprocess
from types import SimpleNamespace

import pytest
from pptx import Presentation
from pptx.util import Inches
from slide_cli import deck


def make_deck(path):
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.shapes.add_textbox(0, 0, Inches(2), Inches(1)).text = "Unicode 目录"
    table = slide.shapes.add_table(1, 2, 0, Inches(1), Inches(4), Inches(1)).table
    table.cell(0, 0).text = "first cell"
    table.cell(0, 1).text = "second cell"
    group = slide.shapes.add_group_shape()
    group.shapes.add_textbox(0, 0, Inches(2), Inches(1)).text = "grouped text"
    presentation.save(path)
    return path


def test_target_ids_order_before_text_and_reload(tmp_path):
    source = make_deck(tmp_path / "目录.pptx")
    snapshot = deck.load_deck_snapshot(source)
    shapes = snapshot.slides[0].shapes
    assert [s.text for s in shapes] == [
        "Unicode 目录",
        "first cell",
        "second cell",
        "grouped text",
    ]
    assert [s.kind for s in shapes][1:3] == ["table_cell", "table_cell"]
    assert shapes[1].parent_target_id == shapes[2].parent_target_id
    patch = deck.DeckPatch(
        "owned",
        [
            deck.TextReplaceOp(1, shapes[0].target_id, shapes[0].text, "updated 目录"),
            deck.TextReplaceOp(1, shapes[1].target_id, "stale", "not applied"),
            deck.TextReplaceOp(1, shapes[3].target_id, shapes[3].text, "group updated"),
        ],
    )
    output = tmp_path / "new.pptx"
    result = deck.apply_deck_patch(source, patch, output_path=output)
    assert result.applied_target_ids == [shapes[0].target_id, shapes[3].target_id]
    assert result.skipped_count == 1
    reloaded = deck.load_deck_snapshot(output).slides[0].shapes
    assert [s.target_id for s in reloaded] == [s.target_id for s in shapes]
    assert [s.text for s in reloaded] == [
        "updated 目录",
        "first cell",
        "second cell",
        "group updated",
    ]
    assert deck.load_deck_snapshot(source) == snapshot
    copied = result.as_dict()
    copied["applied_target_ids"].clear()
    assert result.applied_count == 2


@pytest.mark.parametrize(
    "explicit, environment, expected",
    [
        ("explicit", "environment", "explicit"),
        (None, "environment", "environment"),
        (None, None, "discovered"),
    ],
)
def test_converter_precedence_arguments_and_renamed_output(
    tmp_path, monkeypatch, explicit, environment, expected
):
    source = make_deck(tmp_path / "目录.pptx")
    target = tmp_path / "exports" / "renamed.pdf"
    if environment:
        monkeypatch.setenv("SOFFICE_PATH", environment)
    else:
        monkeypatch.delenv("SOFFICE_PATH", raising=False)
    monkeypatch.setattr(deck.shutil, "which", lambda name: "discovered")
    calls = []

    def convert(command, **kwargs):
        from pathlib import Path

        calls.append((command, kwargs))
        directory = Path(command[command.index("--outdir") + 1])
        (directory / "目录.pdf").write_bytes(b"owned converted output")
        return SimpleNamespace(returncode=0)

    monkeypatch.setattr(deck.subprocess, "run", convert)
    assert (
        deck.export_deck_pdf(
            source, output_path=target, soffice_path=explicit, timeout_sec=19
        )
        == target
    )
    assert target.read_bytes() == b"owned converted output"
    command, kwargs = calls[0]
    assert command[:4] == [expected, "--headless", "--convert-to", "pdf"]
    assert command[-1] == str(source.resolve())
    assert kwargs == dict(
        capture_output=True, text=True, encoding="utf-8", timeout=19, check=False
    )
    assert not (target.parent / "目录.pdf").exists()


@pytest.mark.parametrize("failure", ["nonzero", "timeout", "missing_output"])
def test_conversion_errors_preserve_timing_and_message(tmp_path, monkeypatch, failure):
    source = make_deck(tmp_path / "owned.pptx")

    def fail(command, **kwargs):
        if failure == "timeout":
            raise subprocess.TimeoutExpired(command, kwargs["timeout"])
        return SimpleNamespace(
            returncode=1 if failure == "nonzero" else 0,
            stderr="owned stderr",
            stdout="owned stdout",
        )

    monkeypatch.setattr(deck.subprocess, "run", fail)
    if failure == "timeout":
        with pytest.raises(subprocess.TimeoutExpired) as caught:
            deck.export_deck_pdf(source, soffice_path="owned", timeout_sec=7)
        assert caught.value.timeout == 7
    else:
        message = "owned stderr" if failure == "nonzero" else "expected PDF output"
        with pytest.raises(RuntimeError, match=message):
            deck.export_deck_pdf(source, soffice_path="owned")


def test_missing_source_precedes_converter_lookup(tmp_path, monkeypatch):
    def unexpected(name):
        raise AssertionError("converter lookup before source validation")

    monkeypatch.setattr(deck.shutil, "which", unexpected)
    with pytest.raises(FileNotFoundError):
        deck.export_deck_pdf(tmp_path / "absent.pptx")
