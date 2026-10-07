"""Existing Studio registry through real notebook SQL and format writers."""

import json
from types import SimpleNamespace
from typing import cast

import pytest
from ktem.db.models import Conversation
from ktem.docqa import _runtime_notebook as notebook
from ktem.docqa.artifact_exports import export_artifact_to_path
from ktem.docqa.artifact_generation import (
    artifact_payload_schema,
    build_artifact_payload,
)
from ktem.docqa.artifact_models import SUPPORTED_ARTIFACT_TYPES
from ktem.docqa.artifact_service import build_artifact_note_fields
from ktem.reasoning.mara_artifacts import build_artifact_for_pipeline
from sqlalchemy import Table
from sqlmodel import Session, create_engine

from kotaemon.base import RetrievedDocument


@pytest.fixture
def conversation(tmp_path, monkeypatch):
    engine = create_engine(f"sqlite:///{tmp_path / 'owned-artifacts.db'}")
    cast(Table, Conversation.__table__).create(engine)
    monkeypatch.setattr(notebook, "engine", engine)
    with Session(engine) as session:
        row = Conversation(user="owned-alice", data_source={"sentinel": "retained"})
        session.add(row)
        session.commit()
        session.refresh(row)
        cid = row.id
    try:
        yield engine, cid
    finally:
        engine.dispose()


def generated(artifact_type):
    doc = RetrievedDocument(
        text="Owned observation: the study has 42 participants.",
        doc_id="owned-evidence",
        metadata={
            "file_id": "owned-source",
            "file_name": "owned.txt",
            "page_label": "1",
        },
    )
    return build_artifact_for_pipeline(
        SimpleNamespace(artifact_type=artifact_type, _mara_last_docs=[doc]),
        {"task_type": artifact_type},
    )


@pytest.mark.parametrize("artifact_type", SUPPORTED_ARTIFACT_TYPES)
def test_registry_generation_notebook_export_reload(
    conversation, tmp_path, artifact_type
):
    engine, cid = conversation
    payload = generated(artifact_type)
    assert payload is not None
    assert payload["schema_version"] == "mara_artifact.v1"
    assert set(artifact_payload_schema(artifact_type)["required"]) <= payload.keys()
    assert payload["citations"][0]["source_id"] == "owned-source"
    artifact = notebook.save_captured_artifact(
        cid,
        payload,
        user_id="owned-alice",
        artifact_id="owned-artifact",
        source_scope={"mode": "document", "source_ids": ["owned-source"]},
        timestamp="2026-09-27T00:00:00+00:00",
    )
    assert artifact is not None
    fields = build_artifact_note_fields(artifact)
    note = notebook.save_answer_note_to_conversation(
        cid,
        user_id="owned-alice",
        title=fields["title"],
        answer=fields["text"],
        citation_refs=fields["citation_refs"],
    )
    assert "owned-evidence" in note["citation_refs"]
    formats = ["json", "md", "html"]
    specific = {
        "data_table": "csv",
        "infographic": "svg",
        "slide_deck": "pptx",
        "slide_outline": "pptx",
    }.get(artifact_type)
    if specific:
        formats.append(specific)
    for format_name in formats:
        path = export_artifact_to_path(
            artifact,
            export_format=format_name,
            output_path=tmp_path / f"owned.{format_name}",
        )
        assert path.is_file() and path.stat().st_size > 0
        if format_name == "json":
            assert json.loads(path.read_text(encoding="utf-8")) == artifact
        if format_name == "pptx":
            from pptx import Presentation

            assert len(Presentation(path).slides) > 0
        notebook.record_artifact_export_to_conversation(
            cid,
            "owned-artifact",
            user_id="owned-alice",
            export_format=format_name,
            path=str(path),
        )
    loaded = notebook.get_notebook(cid, user_id="owned-alice")
    record = loaded["artifacts"][0]
    assert record["payload"] == payload
    assert [e["format"] for e in record["exports"]] == formats
    assert record["artifact_id"] == "owned-artifact"
    assert record["created_at"] == "2026-09-27T00:00:00+00:00"
    loaded["artifacts"].clear()
    assert len(notebook.get_notebook(cid, user_id="owned-alice")["artifacts"]) == 1
    with Session(engine) as session:
        row = session.get(Conversation, cid)
        assert row is not None
        assert row.data_source["sentinel"] == "retained"
    notebook.delete_artifact_from_conversation(
        cid, "owned-artifact", user_id="owned-alice"
    )
    assert all((tmp_path / f"owned.{fmt}").exists() for fmt in formats)
    assert notebook.get_notebook(cid, user_id="owned-alice")["notes"][0] == note


def test_private_and_public_read_do_not_grant_notebook_write(conversation):
    engine, cid = conversation
    payload = generated("quiz")
    notebook.save_captured_artifact(
        cid, payload, user_id="owned-alice", artifact_id="q"
    )
    with pytest.raises(notebook.NotebookAccessError):
        notebook.get_notebook(cid, user_id="owned-bob", allow_public=True)
    with Session(engine) as session:
        row = session.get(Conversation, cid)
        assert row is not None
        row.is_public = True
        session.add(row)
        session.commit()
    assert notebook.get_notebook(cid, user_id="owned-bob", allow_public=True)[
        "artifacts"
    ]
    with pytest.raises(notebook.NotebookAccessError):
        notebook.record_artifact_export_to_conversation(
            cid, "q", user_id="owned-bob", export_format="json", path="not-written"
        )
    assert (
        notebook.get_notebook(cid, user_id="owned-alice")["artifacts"][0]["exports"]
        == []
    )


@pytest.mark.parametrize(
    "artifact_type,format_name", [("audio_overview", "mp3"), ("video_overview", "mp4")]
)
def test_media_requires_explicit_adapter(
    conversation, tmp_path, artifact_type, format_name
):
    _engine, cid = conversation
    artifact = notebook.save_captured_artifact(
        cid, generated(artifact_type), user_id="owned-alice"
    )
    assert artifact is not None
    target = tmp_path / f"owned.{format_name}"
    with pytest.raises(ValueError, match="requires a configured media export adapter"):
        export_artifact_to_path(artifact, export_format=format_name, output_path=target)
    assert not target.exists()
    assert artifact["payload"]["media_status"] != "ready"


def test_schema_adapter_failure_never_registers_partial_payload(conversation):
    _engine, cid = conversation
    evidence = generated("quiz")["cited_evidence"]
    with pytest.raises(ValueError, match="missing"):
        build_artifact_payload("quiz", evidence, generation_adapter=lambda req: {})
    assert notebook.get_notebook(cid, user_id="owned-alice")["artifacts"] == []
